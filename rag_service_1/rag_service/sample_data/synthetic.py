"""Synthetic course material in several file formats + a QA set with known answers.

Used by the pytest suite and the notebooks. Facts are distinctive on purpose so retrieval
and grounding can be checked objectively (expected keyword must appear in the answer)."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Dict, List

BIO = {
    "Cell Biology Notes": [
        "Photosynthesis occurs in the chloroplasts of plant cells. Chlorophyll absorbs light energy, which is "
        "used to convert carbon dioxide and water into glucose and oxygen.",
        "Mitochondria are the organelles that produce most of the cell's ATP through cellular respiration. "
        "They have a double membrane and their own circular DNA.",
        "The cell membrane is a phospholipid bilayer. It is selectively permeable and controls which "
        "substances enter and leave the cell. Osmosis is the diffusion of water across this membrane.",
    ],
    "Genetics": [
        "DNA replication is semi-conservative: each new double helix contains one original strand and one "
        "newly synthesised strand. The enzyme DNA polymerase adds nucleotides in the 5' to 3' direction.",
        "Gregor Mendel studied pea plants and formulated the law of segregation and the law of independent "
        "assortment. A dominant allele masks the effect of a recessive allele.",
    ],
}
PHY = {
    "Mechanics": [
        "Newton's second law states that the net force on an object equals its mass multiplied by its "
        "acceleration, F = m a. The SI unit of force is the newton.",
        "Kinetic energy of a moving body is one half of mass times velocity squared. Potential energy near "
        "Earth's surface equals mass times g times height, where g is about 9.8 metres per second squared.",
    ],
    "Optics": [
        "Snell's law relates the angles of incidence and refraction: n1 sin theta1 equals n2 sin theta2. "
        "Total internal reflection occurs when light travels from a denser to a rarer medium beyond the "
        "critical angle.",
        "A convex lens converges parallel light rays to a focus. The focal length is the distance between "
        "the lens and its focal point.",
    ],
}
CS = {
    "Database Normalization": [
        "First normal form requires that every column holds atomic values and there are no repeating groups. "
        "Second normal form removes partial dependencies on a composite primary key.",
        "Third normal form removes transitive dependencies, so non-key attributes depend only on the primary "
        "key. Boyce-Codd normal form is a stricter version of third normal form.",
    ],
    "Transactions": [
        "ACID stands for Atomicity, Consistency, Isolation and Durability. Atomicity means a transaction "
        "either completes fully or has no effect. Durability guarantees committed data survives crashes.",
        "A deadlock occurs when two transactions each wait for a lock held by the other. Databases resolve "
        "deadlocks by aborting one transaction, called the victim.",
    ],
}
SUBJECTS: Dict[str, Dict[str, List[str]]] = {"BIO101": BIO, "PHY101": PHY, "CS201": CS}

# (subject, question, expected keyword in answer, source heading)
QA = [
    ("BIO101", "Where does photosynthesis take place in a plant cell?", "chloroplast", "Cell Biology Notes"),
    ("BIO101", "Which organelle produces ATP?", "mitochondria", "Cell Biology Notes"),
    ("BIO101", "What does semi-conservative DNA replication mean?", "original strand", "Genetics"),
    ("BIO101", "Which enzyme adds nucleotides during DNA replication?", "polymerase", "Genetics"),
    ("PHY101", "State Newton's second law.", "mass", "Mechanics"),
    ("PHY101", "What is the formula for kinetic energy?", "velocity squared", "Mechanics"),
    ("PHY101", "When does total internal reflection occur?", "critical angle", "Optics"),
    ("CS201", "What does ACID stand for in databases?", "durability", "Transactions"),
    ("CS201", "What is third normal form?", "transitive", "Database Normalization"),
    ("CS201", "How are deadlocks resolved?", "victim", "Transactions"),
]
# Questions that must NOT be answered from the material
UNANSWERABLE = [
    ("BIO101", "Who won the football world cup in 1998?"),
    ("PHY101", "What is the recipe for chocolate cake?"),
    ("CS201", "Explain the plot of the film Titanic."),
]
# Cross-subject leakage probes: question about subject X asked while scoped to subject Y
LEAKAGE = [("PHY101", "Which organelle produces ATP?", "mitochondria"), ("BIO101", "State Newton's second law.", "newton")]


def _text(sections: Dict[str, List[str]]) -> str:
    return "\n\n".join(f"{h}\n{p}" for h, ps in sections.items() for p in ps)


def make_pdf(sections: Dict[str, List[str]]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    buf, st = io.BytesIO(), getSampleStyleSheet()
    story = []
    for i, (h, ps) in enumerate(sections.items()):
        story.append(Paragraph(h, st["Heading1"]))
        story += [Paragraph(p, st["BodyText"]) for p in ps]
        if i < len(sections) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def make_docx(sections: Dict[str, List[str]]) -> bytes:
    import docx

    d = docx.Document()
    for h, ps in sections.items():
        d.add_heading(h, level=1)
        for p in ps:
            d.add_paragraph(p)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_pptx(sections: Dict[str, List[str]]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for h, ps in sections.items():
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = h
        s.placeholders[1].text = "\n".join(ps)
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def make_xlsx(sections: Dict[str, List[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for h, ps in sections.items():
        ws = wb.create_sheet(h[:30])
        ws.append(["Topic", "Fact"])
        for p in ps:
            ws.append([h, p])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_all() -> Dict[str, Dict[str, bytes]]:
    """subject -> {filename: bytes}. Different formats per subject on purpose."""
    return {
        "BIO101": {"bio_notes.pdf": make_pdf(BIO), "bio_slides.pptx": make_pptx({"Cell Biology Notes": BIO["Cell Biology Notes"]})},
        "PHY101": {"physics.docx": make_docx(PHY), "physics_facts.xlsx": make_xlsx({"Optics": PHY["Optics"]})},
        "CS201": {
            "db_notes.md": ("# Database Notes\n\n" + _text(CS)).encode(),
            "db_facts.csv": ("topic,fact\n" + "\n".join(f'"{h}","{p}"' for h, ps in CS.items() for p in ps)).encode(),
            "db_page.html": f"<html><head><title>DB</title></head><body><script>x=1</script><p>{_text(CS)}</p></body></html>".encode(),
        },
    }


def write_to_disk(out: Path) -> List[Path]:
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    for subj, files in make_all().items():
        for name, data in files.items():
            p = out / f"{subj}__{name}"
            p.write_bytes(data)
            paths.append(p)
    return paths


if __name__ == "__main__":
    for p in write_to_disk(Path(__file__).parent / "files"):
        print(p)
